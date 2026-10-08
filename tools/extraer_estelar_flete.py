# -*- coding: utf-8 -*-
"""Extrae tarifas Estelar (COP/kg) y pesos de FIBRA OPTICA del Excel SAT."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "assets" / "files" / "SAT SEGURIDAD ELECTRONICA V3 2026.xlsx"
OUT_JSON = ROOT / "assets" / "files" / "estelar_flete.json"
OUT_JS = ROOT / "assets" / "files" / "estelar_flete.data.js"
JSON_KEY = "assets/files/estelar_flete.json"

ESTELAR_EXCEPT_SKUS = ["7655", "7656", "7654", "7661", "7660", "14181", "15732", "9939"]


def sku_key(v) -> str:
    s = str(v or "").strip()
    if not s or s.lower() == "nan":
        return ""
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".", 1)[0]
    return s


def num(v) -> float:
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(" ", "").replace(".", "").replace(",", ".")
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        try:
            return float(s)
        except (TypeError, ValueError):
            return 0.0


def main() -> None:
    wb = load_workbook(XLSX, data_only=True, read_only=True)
    cities = []
    ws_c = wb["Ciudades Flete"]
    for i, row in enumerate(ws_c.iter_rows(min_row=2, values_only=True), start=2):
        name = str(row[0] or "").strip()
        if not name:
            continue
        rate = num(row[2])
        dept = str(row[1] or "").strip()
        cities.append({"name": name, "departamento": dept, "valor_kilo": round(rate, 6)})
    print("ciudades", len(cities), "first", cities[0], "armenia", [c for c in cities if "ARMENIA" in c["name"].upper()][:3])

    products = {}
    ws_f = wb["FIBRA OPTICA"]
    for row in ws_f.iter_rows(min_row=5, values_only=True):
        sku = sku_key(row[1] if len(row) > 1 else "")
        if not sku:
            continue
        kg = num(row[8] if len(row) > 8 else 0)
        usd = num(row[4] if len(row) > 4 else 0)
        cop_uni = num(row[5] if len(row) > 5 else 0)
        nearest = round(usd)
        if abs(usd - nearest) <= 0.02:
            usd = float(nearest)
        else:
            usd = round(usd, 2)
        products[sku] = {
            "kg": kg,
            "usd": usd,
            "cop_uni": round(cop_uni, 6),
            "name": str(row[2] or "").strip(),
        }
    print("productos fibra", products)

    payload = {
        "transportadora": "estelar",
        "valor_sobre_producto": 0.005,
        "except_skus": ESTELAR_EXCEPT_SKUS,
        "sat_precio_skus": ESTELAR_EXCEPT_SKUS,
        "ciudades": cities,
        "productos": products,
    }
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    js = (
        "window.__IZC_DATA__=window.__IZC_DATA__||{};\n"
        f"window.__IZC_DATA__[{json.dumps(JSON_KEY)}]={json.dumps(payload, ensure_ascii=False)};\n"
    )
    OUT_JS.write_text(js, encoding="utf-8")
    print("wrote", OUT_JSON, "and", OUT_JS.name)


if __name__ == "__main__":
    main()
